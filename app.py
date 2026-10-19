import json
import os
import re
from io import BytesIO

from flask import Flask, request, send_file, jsonify, render_template_string
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

app = Flask(__name__)

MODEL = "gemini-3.8-flash"

THEMES = {
    "ko'k": {"bg": "1F3A5F", "title": "FFFFFF", "text": "E8EEF5"},
    "yashil": {"bg": "1E5B4F", "title": "FFFFFF", "text": "E6F2EF"},
    "qora": {"bg": "1A1A1A", "title": "FFD166", "text": "EEEEEE"},
    "oq": {"bg": "FFFFFF", "title": "1F3A5F", "text": "333333"},
}

LANGUAGES = {"uz": "o'zbek tilida (lotin yozuvi)", "ru": "rus tilida", "en": "ingliz tilida"}


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def paint_background(slide, hex_color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(hex_color)


# ---------- AI: mavzudan slaydlar matnini yaratish (Gemini) ----------
def generate_content(topic, count, lang_key):
    """Gemini'dan {"title": ..., "slides": [{"title":..., "bullets":[...]}]} oladi."""
    if not os.environ.get("GEMINI_API_KEY"):
        raise RuntimeError(
            "GEMINI_API_KEY topilmadi. Terminalda: export GEMINI_API_KEY='kalitingiz'"
        )
    from google import genai  # pip install google-genai

    language = LANGUAGES.get(lang_key, LANGUAGES["uz"])
    prompt = f"""Quyidagi mavzu bo'yicha taqdimot (slaydlar) tayyorla.

Mavzu: {topic}
Kontent slaydlari soni: {count} ta (sarlavha slaydini hisobga olmagan holda)
Til: {language}

Talablar:
- Har bir slaydda sarlavha va 3-5 ta qisqa, aniq band bo'lsin (har biri bir jumla, 15 so'zdan oshmasin).
- Slaydlar mantiqiy ketma-ketlikda bo'lsin: kirish, asosiy qismlar, xulosa.
- Faqat tekshirib bo'ladigan, aniq ma'lumot yoz; ishonchsiz raqam va faktlar o'ylab topma.

Javobni FAQAT JSON ko'rinishida qaytar, boshqa hech narsa yozma:
{{"title": "taqdimot nomi", "slides": [{{"title": "slayd sarlavhasi", "bullets": ["band 1", "band 2"]}}]}}"""

    client = genai.Client()  # kalitni GEMINI_API_KEY o'zgaruvchisidan oladi
    interaction = client.interactions.create(model=MODEL, input=prompt)
    text = (interaction.output_text or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text)
    start, end = text.find("{"), text.rfind("}")
    if start == -1 or end == -1:
        raise ValueError("AI javobi JSON emas, qayta urinib ko'ring")
    data = json.loads(text[start : end + 1])
    if not data.get("slides"):
        raise ValueError("AI slaydlarni qaytarmadi, qayta urinib ko'ring")
    return data


# ---------- PowerPoint fayl yaratish ----------
def build_presentation(deck_title, slides, theme_name):
    theme = THEMES.get(theme_name, THEMES["ko'k"])
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Sarlavha slaydi
    s = prs.slides.add_slide(prs.slide_layouts[0])
    paint_background(s, theme["bg"])
    s.shapes.title.text = deck_title or "Taqdimot"
    for p in s.shapes.title.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(48)
            r.font.bold = True
            r.font.color.rgb = rgb(theme["title"])
    s.placeholders[1].text = ""

    # Kontent slaydlari
    for item in slides:
        title = (item.get("title") or "").strip()
        bullets = [str(b).strip() for b in (item.get("bullets") or []) if str(b).strip()]
        if not title and not bullets:
            continue
        s = prs.slides.add_slide(prs.slide_layouts[1])
        paint_background(s, theme["bg"])

        s.shapes.title.text = title
        s.shapes.title.left, s.shapes.title.top = Inches(0.8), Inches(0.5)
        s.shapes.title.width, s.shapes.title.height = Inches(11.7), Inches(1.3)
        for p in s.shapes.title.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(38)
                r.font.bold = True
                r.font.color.rgb = rgb(theme["title"])

        body = s.placeholders[1]
        body.left, body.top = Inches(0.8), Inches(2.0)
        body.width, body.height = Inches(11.7), Inches(4.9)
        tf = body.text_frame
        tf.clear()
        tf.word_wrap = True
        for i, line in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.space_after = Pt(12)
            for r in p.runs:
                r.font.size = Pt(26)
                r.font.color.rgb = rgb(theme["text"])

    buf = BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


@app.route("/")
def index():
    return render_template_string(PAGE, themes=list(THEMES.keys()))


@app.route("/generate", methods=["POST"])
def generate():
    data = request.get_json(silent=True) or {}
    topic = (data.get("topic") or "").strip()
    if not topic:
        return jsonify(error="Mavzuni kiriting"), 400
    try:
        count = max(3, min(15, int(data.get("count", 7))))
    except (TypeError, ValueError):
        count = 7
    try:
        content = generate_content(topic, count, data.get("lang", "uz"))
        buf = build_presentation(
            content.get("title") or topic, content["slides"], data.get("theme", "ko'k")
        )
    except Exception as e:
        return jsonify(error=str(e)), 500
    return send_file(
        buf,
        as_attachment=True,
        download_name="taqdimot.pptx",
        mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )


PAGE = """<!doctype html>
<html lang="uz">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Slayd yaratuvchi</title>
<style>
  *{box-sizing:border-box}
  body{font-family:system-ui,Arial,sans-serif;background:#f3f5f9;margin:0;color:#222}
  .wrap{max-width:640px;margin:0 auto;padding:40px 16px 60px}
  h1{margin:0 0 4px}
  .sub{color:#667;margin-bottom:24px}
  .card{background:#fff;border-radius:12px;padding:20px;box-shadow:0 1px 4px rgba(0,0,0,.06)}
  label{font-weight:600;font-size:14px;display:block;margin:14px 0 6px}
  label:first-child{margin-top:0}
  input,select{width:100%;padding:11px 12px;border:1px solid #cfd6e2;border-radius:8px;font:inherit;background:#fff}
  .row{display:grid;grid-template-columns:1fr 1fr 1fr;gap:12px}
  @media(max-width:560px){.row{grid-template-columns:1fr}}
  button{border:0;border-radius:8px;font:inherit;font-weight:600;cursor:pointer}
  .go{background:#1F3A5F;color:#fff;width:100%;padding:14px;font-size:17px;margin-top:20px}
  .go:disabled{opacity:.6;cursor:wait}
  #msg{margin-top:12px;min-height:20px;font-size:14px}
  .err{color:#b42318}.ok{color:#1e6b3a}
</style>
</head>
<body>
<div class="wrap">
  <h1>Slayd yaratuvchi</h1>
  <div class="sub">Mavzuni yozing, slaydlarni sun'iy intellekt o'zi tayyorlaydi.</div>
  <div class="card">
    <label for="topic">Taqdimot mavzusi</label>
    <input id="topic" placeholder="Masalan: Quyosh energiyasining kelajagi" autofocus>

    <div class="row">
      <div>
        <label for="count">Slaydlar soni</label>
        <select id="count">
          <option>5</option><option selected>7</option><option>10</option><option>12</option><option>15</option>
        </select>
      </div>
      <div>
        <label for="lang">Til</label>
        <select id="lang">
          <option value="uz">O'zbekcha</option>
          <option value="ru">Русский</option>
          <option value="en">English</option>
        </select>
      </div>
      <div>
        <label for="theme">Rang</label>
        <select id="theme">
          {% for t in themes %}<option value="{{t}}">{{t}}</option>{% endfor %}
        </select>
      </div>
    </div>

    <button class="go" id="go" onclick="generate()">Taqdimotni yaratish</button>
    <div id="msg"></div>
  </div>
</div>

<script>
document.getElementById('topic').addEventListener('keydown',e=>{if(e.key==='Enter')generate();});
async function generate(){
  const msg=document.getElementById('msg'), btn=document.getElementById('go');
  const topic=document.getElementById('topic').value.trim();
  msg.className=''; msg.textContent='';
  if(!topic){msg.className='err';msg.textContent='Mavzuni kiriting';return;}
  btn.disabled=true; btn.textContent='Slaydlar tayyorlanmoqda (20-40 soniya)...';
  try{
    const r=await fetch('/generate',{method:'POST',headers:{'Content-Type':'application/json'},
      body:JSON.stringify({topic,
        count:document.getElementById('count').value,
        lang:document.getElementById('lang').value,
        theme:document.getElementById('theme').value})});
    if(!r.ok){const e=await r.json();throw new Error(e.error||'Xatolik');}
    const blob=await r.blob();
    const a=document.createElement('a');
    a.href=URL.createObjectURL(blob); a.download='taqdimot.pptx'; a.click();
    msg.className='ok'; msg.textContent='Tayyor! Fayl yuklab olindi.';
  }catch(e){msg.className='err';msg.textContent=e.message;}
  btn.disabled=false; btn.textContent='Taqdimotni yaratish';
}
</script>
</body>
</html>
"""

if __name__ == "__main__":
    # Lokal ishga tushirish. Serverda gunicorn ishlatiladi (README ga qarang).
    # Debug faqat FLASK_DEBUG=1 bo'lganda yoqiladi.
    app.run(
        host=os.environ.get("HOST", "127.0.0.1"),
        port=int(os.environ.get("PORT", 5000)),
        debug=os.environ.get("FLASK_DEBUG") == "1",
    )
